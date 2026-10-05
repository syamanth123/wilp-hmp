"""
Reference prototype for the Course Dibba importer.

Reads the two real schedule files and writes normalized CSVs. Claude Code
should port THESE RULES to TypeScript (packages/integrations/src/dibba/*).
The TypeScript version is the one that ships; this file is the spec-by-example.

Usage:
  python parse_dibba.py <dibba.docx> <dibba.xlsx> <out_dir>
(.doc must first be converted: soffice --headless --convert-to docx file.doc)
"""
import re, sys, csv
from collections import Counter
import docx, openpyxl

# ---- Course-code normalisation (extends packages/db/src/course-code.ts) ----
# Real data needs two extras the current normaliser rejects:
#   * a trailing "T" for dissertation / project courses  (SS ZG628T, BITS ZC425T)
#   * 5-letter prefixes                                  (POWAB ZC113)
CODE_RE = re.compile(r'\b([A-Z]{2,5})\s*(Z\s*[CG])\s*(\d)\s*(\d)\s*(\d)(\d?)(T?)\b')
def find_codes(text):
    out = []
    for m in CODE_RE.finditer(text.upper()):
        prefix, zc = m.group(1), m.group(2).replace(' ', '')
        digits = m.group(3) + m.group(4) + m.group(5) + m.group(6)
        out.append((f"{prefix} {zc}{digits}{m.group(7)}", m.start(), m.end()))
    return out

# ---- Slot header -> (slot_no, day, session) ----
DAY = {'SAT': 'SAT', 'SUN': 'SUN', 'FRI': 'FRI'}
SESSION = {'FN': 'FN', 'AN': 'AN', 'EV': 'EV', 'EN': 'EV'}   # EN is a typo-variant of EV
CANON = {('SAT','FN'):1, ('SAT','AN'):2, ('SUN','FN'):3, ('SUN','AN'):4,
         ('FRI','FN'):5, ('FRI','AN'):6, ('SAT','EV'):7, ('SUN','EV'):8}
def parse_slot_header(h, col_index):
    t = h.upper().replace('(', ' ').replace(')', ' ')
    sl = re.search(r'SL\s*(\d)', t)
    day = next((DAY[w] for w in re.findall(r'[A-Z]+', t) if w in DAY), None)
    ses = next((SESSION[w] for w in re.findall(r'[A-Z]+', t) if w in SESSION), None)
    slot = int(sl.group(1)) if sl else CANON.get((day, ses))
    warn = ''
    if slot is None:
        slot = col_index  # last resort: column position (col 1 == SL1)
        warn = f'slot inferred from column position for header "{h.strip()}"'
    elif sl and (day, ses) in CANON and CANON[(day, ses)] != slot:
        warn = f'header "{h.strip()}" disagrees with standard slot map (SL{slot} vs {day} {ses})'
    return slot, day, ses, warn

BATCH_RE = re.compile(r'([12])\s*/\s*(20\d\d)')
TIME_RE = re.compile(r'\b(SAT(?:URDAY)?|SUN(?:DAY)?|FRI(?:DAY)?)\s*(\d{1,2})[.:](\d{2})', re.I)

def parse_batch_cell(txt):
    m = BATCH_RE.search(txt)
    # Student count: prefer a bracketed number "(178)"; "(00)" means zero.
    # 4-digit 50xx-51xx numbers are ERP term codes ("2/2020 (5092)", "(5105)"), NOT head-counts.
    # Fallback: a number alone on its own line ("2/2024\n502").
    # Numbers followed by words ("4 Core", "2 EL") are course-mix notes, ignored.
    rest = BATCH_RE.sub(' ', txt)
    br = [n for n in re.findall(r'\((\d{1,4})\)', rest) if not re.fullmatch(r'5[01]\d\d', n)]
    bare = [n for n in re.findall(r'(?m)^\s*(\d{1,4})\s*$', rest) if not re.fullmatch(r'5[01]\d\d', n)]
    pick = (br or bare or [''])[-1]
    count = type('M', (), {'group': lambda self, i: str(int(pick)) if pick else ''})()
    return {
        'admit_batch': f"{m.group(1)}/{m.group(2)}" if m else '',
        'is_new_admission': 'NEW ADM' in txt.upper(),
        'is_backlog_row': 'BACKLOG' in txt.upper(),
        'student_count': count.group(1) if count else '',
    }

PROG_RE = re.compile(r'^\s*((?:[A-Z]{2}\s?\d{2}|\d{2}[A-Z]{2})(?:\s*/\s*(?:[A-Z]{2}\s?\d{2}|\d{2}[A-Z]{2}))*)')
def programme_code(title):
    # "HB28 MBA (...)" -> "HB28"; "18BT/18ET B.Tech." -> "18BT/18ET"; "HT 31 M.Tech." -> "HT31"
    m = PROG_RE.match(title.upper())
    return re.sub(r'\s+', '', m.group(1)) if m else ''

def classify(fragment):
    u = fragment.upper()
    if 'NOT OFFERED' in u: return 'NOT_OFFERED'
    if 'BACKLOG' in u or 'BKLG' in u: return 'BACKLOG'
    if 'CORE' in u: return 'CORE'
    return 'ELECTIVE_OR_UNSPECIFIED'

def parse_doc(path):
    d = docx.Document(path)
    rows, warnings = [], []
    programme = ''
    for ti, t in enumerate(d.tables):
        # NOTE: table 0 holds the INDEX *and* the first programme (18BT/18ET);
        # rows are only consumed after an "Admit Batch" header row is seen.
        header = None
        for r in t.rows:
            cells = []
            for c in r.cells:  # de-duplicate merged cells
                if not cells or c._tc is not cells[-1][0]:
                    cells.append((c._tc, c.text))
            texts = [x[1] for x in cells]
            if len(texts) == 1 or (texts and not header and not texts[0].strip().lower().startswith('admit')):
                if texts[0].strip():
                    programme = re.sub(r'\s+', ' ', texts[0]).strip()
                continue
            if texts[0].strip().lower().startswith('admit'):
                header = [(t_, *parse_slot_header(t_, i)) for i, t_ in enumerate(texts)]
                for h in header[1:]:
                    if h[4]: warnings.append(f"table {ti} {programme[:40]}: {h[4]}")
                continue
            if header is None:
                continue
            batch = parse_batch_cell(texts[0])
            for ci, cell in enumerate(texts[1:], 1):
                if ci >= len(header) or not cell.strip():
                    continue
                _, slot, day, ses, _w = header[ci]
                codes = find_codes(cell)
                if not codes:
                    warnings.append(f"table {ti} {programme[:30]} batch {batch['admit_batch']} SL{slot}: no course code in '{cell.strip()[:60]}'")
                for k, (code, s, e) in enumerate(codes):
                    nxt = codes[k + 1][1] if k + 1 < len(codes) else len(cell)
                    frag = cell[s:nxt]
                    title = re.sub(r'^[\s|:]+', '', cell[e:nxt]).split('\n')[0]
                    title = re.sub(r'\(.*$', '', title).strip(' /|:')
                    tm = TIME_RE.search(frag)
                    rows.append({
                        'source': 'doc', 'programme': programme,
                        'programme_code': programme_code(programme),
                        **batch, 'slot': slot, 'slot_day': day or '', 'slot_session': ses or '',
                        'course_code': code, 'course_title': title,
                        'course_type': classify(frag),
                        'class_time_hint': f"{tm.group(1)[:3].upper()} {tm.group(2)}:{tm.group(3)}" if tm else '',
                        'raw_cell': re.sub(r'\s+', ' ', frag).strip()[:120],
                    })
    return rows, warnings

def parse_xlsx(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb['Course Dibba S1-2024']
    hdr = [c.value for c in ws[1]]
    ix = {h: i for i, h in enumerate(hdr) if h}
    dibba = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r[ix['Subject']]: continue
        found = find_codes(f"{r[ix['Subject']]}{r[ix['Catalog']]}")
        if not found: continue
        dibba.append({
            'acad_plan': r[ix['Acad Plan']], 'degree': (r[ix['Degree Programme']] or '').strip(),
            'programme': r[ix['Programme']], 'admit_batch': str(r[ix['Admit Batch']]).replace('|', '/'),
            'degree_semester': r[ix['Degree Semester']], 'student_count': r[ix['Active Student No.']],
            'domain': r[ix['Domain']], 'course_type': str(r[ix['Type']]).upper(),
            'slot': r[ix['Exam Slot']], 'erp_course_id': r[ix['Course ID']],
            'course_code': found[0][0], 'course_title': r[ix['Unique Title']] or r[ix['Descr']],
            'min_units': r[ix['Min Units']], 'remarks': r[ix['Remarks']] or '',
        })
    fws = wb['Sheet3']
    fh = [c.value for c in fws[1]]
    fx = {h: i for i, h in enumerate(fh) if h}
    faculty = []
    for r in fws.iter_rows(min_row=2, values_only=True):
        if not r[fx['Course Number']]: continue
        found = find_codes(str(r[fx['Course Number']]))
        if not found: continue
        name = str(r[fx['FACULTY NAME']] or '').strip()
        faculty.append({
            'course_code': found[0][0], 'slot': r[fx['dabba ']],
            'faculty_name': re.sub(r'\s*\(LEAD\)\s*', '', name, flags=re.I),
            'is_lead': '(LEAD)' in name.upper(),
            'email': (r[fx['Email']] or '').strip().lower(),
            'psrn_or_gfid': r[fx['PSRN/GFID']] or '', 'department': r[fx['Department']] or '',
            'campus': r[fx['campus']] or '',
        })
    return dibba, faculty

def write(path, rows):
    if not rows: return
    with open(path, 'w', newline='', encoding='utf-8') as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0].keys())); w.writeheader(); w.writerows(rows)

if __name__ == '__main__':
    docp, xlsxp, out = sys.argv[1:4]
    doc_rows, warns = parse_doc(docp)
    x_rows, fac = parse_xlsx(xlsxp)
    write(f'{out}/dibba_S1_2025-26_from_doc.csv', doc_rows)
    write(f'{out}/dibba_S1_2024_from_xlsx.csv', x_rows)
    write(f'{out}/faculty_course_map_2024.csv', fac)
    with open(f'{out}/parse_warnings.txt', 'w') as f: f.write('\n'.join(warns))
    print('doc rows', len(doc_rows), '| unique courses', len({r['course_code'] for r in doc_rows}),
          '| programmes', len({r['programme'] for r in doc_rows}), '| warnings', len(warns))
    print('xlsx rows', len(x_rows), '| faculty rows', len(fac), '| faculty with email', sum(1 for f in fac if f['email']))
    print('slots', Counter(r['slot'] for r in doc_rows))
    print('types', Counter(r['course_type'] for r in doc_rows))
    dc = {r['course_code'] for r in doc_rows}; fc = {f['course_code'] for f in fac if f['email']}
    print('2025 courses with a 2024 faculty match', len(dc & fc), 'of', len(dc))
